"""AnonymAIzer: Streamlit frontend with local/Hugging Face model loading."""
from __future__ import annotations

import os
import json
import re
from io import BytesIO
from xml.sax.saxutils import escape
from collections import Counter, defaultdict
from pathlib import Path

import streamlit as st
from utils.auth import require_login, auth_enabled

ROOT = Path(__file__).resolve().parent
MODEL_DIR = Path(os.environ.get('ANONYM_AIZER_MODEL', r'C:\AI\models\rizzo-pii-0.3B'))
HF_MODEL = "rizzoaiacademy/rizzo-pii-0.3B"
HF_REVISION = "v1.5.0"
USE_LOCAL_MODEL = MODEL_DIR.is_dir() and (MODEL_DIR / "config.json").is_file()
MODEL_SOURCE = str(MODEL_DIR) if USE_LOCAL_MODEL else HF_MODEL
LABELS = {
    'FULLNAME': 'Nomi e cognomi', 'AGE': 'Età', 'GENDER': 'Genere',
    'DATE': 'Date', 'TIME': 'Orari', 'STREET': 'Strade',
    'BUILDINGNUM': 'Numeri civici', 'ZIPCODE': 'CAP', 'CITY': 'Città',
    'PROVINCE': 'Province', 'EMAIL': 'Email', 'TELEPHONENUM': 'Telefoni',
    'CF': 'Codici fiscali', 'PIVA': 'Partite IVA', 'ID_DOC': 'Documenti identificativi',
    'IBAN': 'IBAN', 'CREDITCARDNUMBER': 'Carte di credito',
    'AMOUNT': 'Importi', 'TARGA': 'Targhe', 'ORG': 'Organizzazioni',
    'DOCID': 'Codici di atti', 'CATASTO': 'Dati catastali',
    'URL': 'URL', 'IPADDR': 'Indirizzi IP',
}
DETECTORS = {
    'EMAIL': re.compile(r'(?<![\w.@])[\w.+%-]+@[\w.-]+\.[A-Za-z]{2,}(?![\w.-])'),
    'URL': re.compile(r'\bhttps?://[^\s<>"\']+', re.I),
    'IPADDR': re.compile(r'(?<!\d)(?:\d{1,3}\.){3}\d{1,3}(?!\d)'),
    'IBAN': re.compile(r'\bIT\d{2}[A-Z]\d{10}[A-Z0-9]{12}\b', re.I),
    'CF': re.compile(r'\b[A-Z]{6}\d{2}[A-EHLMPRST][0-7]\d[A-Z]\d{3}[A-Z]\b', re.I),
}

favicon = ROOT / 'assets' / 'favicon.png'
st.set_page_config(page_title='AnonymAIzer | Anonimizzazione locale',
                   page_icon=str(favicon) if favicon.is_file() else '🔒', layout='wide')

# Proteggi la pagina PRIMA di leggere documenti o caricare il modello.
user_email = require_login()
st.markdown('''<style>
/* AnonymAIzer: palette blu, anche per le etichette selezionate. */
:root { --anony-blue: #1769aa; --anony-blue-dark: #115486; }
.block-container{padding-top:2rem;max-width:1200px}
[data-testid="stMetric"]{border:1px solid #d9e4ee;border-radius:12px;padding:12px}
/* Elementi selezionati nel multiselect della sidebar. */
[data-baseweb="tag"] { background-color: #dceefa !important; color: #145a91 !important; }
[data-baseweb="tag"] span { color: #145a91 !important; }
[data-baseweb="tag"] [role="button"], [data-baseweb="tag"] svg {
  color: #145a91 !important; fill: #145a91 !important;
}
/* Comandi primari e indicatori di selezione. */
button[kind="primary"] { background: var(--anony-blue) !important; border-color: var(--anony-blue) !important; }
button[kind="primary"]:hover { background: var(--anony-blue-dark) !important; border-color: var(--anony-blue-dark) !important; }
input[type="radio"], input[type="checkbox"] { accent-color: var(--anony-blue); }
</style>''', unsafe_allow_html=True)

@st.cache_resource(show_spinner='Caricamento di Rizzo-PII in corso...')
def load_model(model_path: str, revision: str | None):
    from transformers import pipeline
    # Local path on Windows; download once into the container cache on Cloud.
    return pipeline(task="token-classification", model=model_path,
                    tokenizer=model_path, revision=revision,
                    aggregation_strategy="simple", device="cpu")


def chunks(text: str, max_words: int = 105, overlap: int = 15):
    """Yield character spans of overlapping word chunks."""
    words = list(re.finditer(r'\S+', text))
    if not words:
        return
    stride = max(1, max_words - overlap)
    for i in range(0, len(words), stride):
        group = words[i:i + max_words]
        if not group:
            break
        left, right = group[0].start(), group[-1].end()
        yield left, text[left:right]
        if i + max_words >= len(words):
            break


def normalize_label(raw: str) -> str:
    label = str(raw).upper().replace('B-', '').replace('I-', '')
    return {'PHONE': 'TELEPHONENUM', 'PERSON': 'FULLNAME',
            'IP': 'IPADDR', 'CREDIT_CARD': 'CREDITCARDNUMBER'}.get(label, label)


def find_entities(text: str, nlp):
    spans = []
    for offset, chunk in chunks(text):
        for e in nlp(chunk):
            label = normalize_label(e.get('entity_group', e.get('entity', 'UNKNOWN')))
            a, b = offset + int(e['start']), offset + int(e['end'])
            if a < b and text[a:b].strip():
                spans.append({'start': a, 'end': b, 'label': label,
                              'score': float(e.get('score', 0)), 'source': 'Modello'})
    for label, regex in DETECTORS.items():
        for m in regex.finditer(text):
            a, b = m.span()
            if label == 'IPADDR' and any(int(x) > 255 for x in m.group().split('.')):
                continue
            spans.append({'start': a, 'end': b, 'label': label,
                          'score': None, 'source': 'Regola'})
    # Exact duplicates occur at overlapping chunk boundaries. Prefer rules for identifiers,
    # then confidence, then longer candidates, and do not replace overlapping spans twice.
    spans.sort(key=lambda x: (0 if x['source'] == 'Regola' else 1,
                              -(x['score'] or 0), -(x['end'] - x['start'])))
    chosen = []
    for s in spans:
        if not any(s['start'] < t['end'] and t['start'] < s['end'] for t in chosen):
            chosen.append(s)
    return sorted(chosen, key=lambda x: x['start'])


def anonymize(text: str, spans: list, enabled: set):
    counters = defaultdict(int)
    assigned = {}
    display_rows = []
    output = []
    cursor = 0
    for s in spans:
        a, b, label = s['start'], s['end'], s['label']
        value = text[a:b]
        masked = label in enabled
        output.append(text[cursor:a])
        if masked:
            key = (label, value.casefold())
            if key not in assigned:
                counters[label] += 1
                assigned[key] = f'[{label}_{counters[label]}]'
            placeholder = assigned[key]
            output.append(placeholder)
        else:
            placeholder = ''
            output.append(value)
        display_rows.append({'Categoria': LABELS.get(label, label), 'Valore originale': value,
                             'Segnaposto': placeholder if masked else '—',
                             'Rilevazione': s['source'],
                             'Confidenza': round(s['score'], 3) if s['score'] is not None else '—'})
        cursor = b
    output.append(text[cursor:])
    return ''.join(output), display_rows, sum(bool(r['Segnaposto'] != '—') for r in display_rows)




def read_document(upload):
    """Extract readable text; no OCR and no preservation of page layout."""
    extension = Path(upload.name).suffix.lower()
    data = upload.getvalue()
    if extension == '.docx':
        from docx import Document
        doc = Document(BytesIO(data))
        paragraphs = []
        # Include normal paragraphs and cells in their actual body order.
        from docx.oxml.ns import qn
        from docx.text.paragraph import Paragraph
        from docx.table import Table
        for child in doc.element.body.iterchildren():
            if child.tag == qn('w:p'):
                txt = Paragraph(child, doc).text
                if txt.strip():
                    paragraphs.append(txt)
            elif child.tag == qn('w:tbl'):
                table = Table(child, doc)
                for row in table.rows:
                    paragraphs.append(' | '.join(cell.text.replace('\n', ' / ') for cell in row.cells))
        return '\n\n'.join(paragraphs)
    if extension == '.pdf':
        import fitz
        with fitz.open(stream=data, filetype='pdf') as pdf:
            return '\n\n'.join(page.get_text('text', sort=True) for page in pdf)
    raise ValueError('Formato non supportato: seleziona DOCX o PDF.')


def make_docx(text: str):
    from docx import Document
    document = Document()
    for paragraph in text.split('\n'):
        document.add_paragraph(paragraph)
    buff = BytesIO()
    document.save(buff)
    return buff.getvalue()


def make_pdf(text: str):
    # This generates a new PDF; it does not edit/redact the source PDF in place.
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, KeepTogether
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.lib.styles import ParagraphStyle
    buffer = BytesIO()
    font_name = 'Helvetica'
    # Prefer a Unicode font available on Windows for accented Italian characters.
    for font_path in (r'C:\Windows\Fonts\arial.ttf', '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'):
        if Path(font_path).is_file():
            font_name = 'AnonymAIzerUnicode'
            if font_name not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(TTFont(font_name, font_path))
            break
    style = ParagraphStyle('BodyA', fontName=font_name, fontSize=10, leading=15, spaceAfter=6, alignment=TA_LEFT, wordWrap='CJK')
    doc = SimpleDocTemplate(buffer, pagesize=A4, leftMargin=42, rightMargin=42, topMargin=45, bottomMargin=45)
    story = []
    for line in text.split('\n'):
        story.append(Paragraph(escape(line) if line.strip() else '&nbsp;', style))
    doc.build(story)
    return buffer.getvalue()

if 'output' not in st.session_state:
    st.session_state.output = ''
    st.session_state.rows = []
    st.session_state.count = 0
    st.session_state.source_name = 'testo'

logo = ROOT / 'assets' / 'logo.png'
if logo.exists():
    st.image(str(logo), width=380)
st.caption('Anonimizza testi e proteggi i dati personali con AI. I testi sono elaborati sul server che ospita l’app.')

# Il modello viene caricato dopo il login; la cache e condivisa tra sessioni
# nello stesso processo e sopravvive ai rerun Streamlit.
nlp = None
model_error = None
try:
    with st.spinner('Caricamento iniziale di Rizzo-PII (il primo avvio cloud può richiedere tempo)...'):
        nlp = load_model(MODEL_SOURCE, None if USE_LOCAL_MODEL else HF_REVISION)
except Exception as exc:
    model_error = f'{type(exc).__name__}: {exc}'
    st.error(f'Impossibile caricare Rizzo-PII: {model_error}')

with st.sidebar:
    if auth_enabled():
        st.caption(f'Accesso: {user_email}')
        if st.button('Esci', key='logout_anonymaizer'):
            st.logout()
    st.header('Impostazioni')
    st.subheader('Categorie da anonimizzare')
    all_tags = list(LABELS)
    # Include anche eventuali nuove etichette presenti nella configurazione del modello.
    try:
        model_labels = nlp.model.config.id2label.values() if nlp is not None else []
        for raw_label in model_labels:
            tag = normalize_label(raw_label)
            if tag not in all_tags and tag not in ('O', 'UNKNOWN', 'LABEL_0') and not tag.startswith('LABEL_'):
                all_tags.append(tag)
    except (OSError, ValueError, TypeError):
        pass
    selected = st.multiselect('Seleziona le categorie', options=all_tags,
                              default=all_tags,
                              format_func=lambda tag: f'{LABELS.get(tag, tag)} ({tag})',
                              help='All’apertura tutte le categorie sono selezionate, comprese città e organizzazioni.')
    st.info('Su Streamlit Community Cloud i documenti vengono inviati al server cloud prima dell’anonimizzazione. Per questa fase usa solo testi sintetici.')
    st.caption(f'Modello Rizzo-PII 0.3B · CPU · {"Disco locale" if USE_LOCAL_MODEL else "Hugging Face / cache cloud"}')
    try:
        import psutil
        mem_mb = psutil.Process().memory_info().rss / (1024 ** 2)
        st.caption(f'Memoria processo: {mem_mb:,.0f} MB (RSS)')
    except Exception:
        pass

mode = st.radio('Modalità di inserimento', ['Incolla testo', 'Carica documento'], horizontal=True)
left, right = st.columns(2, gap='large')
with left:
    st.subheader('Testo originale' if mode == 'Incolla testo' else 'Documento originale')
    upload = None
    input_text = ''
    if mode == 'Incolla testo':
        st.caption('Incolla il testo da anonimizzare')
        input_text = st.text_area('Incolla il testo da anonimizzare', height=330,
                                  label_visibility='collapsed',
                                  placeholder='Es.: Mario Rossi, email mario.rossi@example.com ...',
                                  key='input_text')
    else:
        upload = st.file_uploader('Carica un documento Word o PDF', type=['docx', 'pdf'],
                                  help='DOCX e PDF testuali. I PDF scannerizzati richiedono OCR (non incluso).')
        if upload is not None:
            try:
                input_text = read_document(upload)
                if not input_text.strip():
                    st.warning('Nessun testo estraibile. Il PDF potrebbe essere scannerizzato: serve OCR.')
                else:
                    st.caption(f'{len(input_text):,} caratteri estratti da {upload.name}')
                    with st.expander('Anteprima del testo estratto'):
                        st.text_area('Testo estratto', value=input_text[:12000], height=220, disabled=True)
            except Exception as exc:
                st.error(f'Impossibile leggere il documento: {type(exc).__name__}: {exc}')
    btn_col, reset_col = st.columns([2, 1])
    clicked = btn_col.button('Anonimizza', type='primary', use_container_width=True)
    if reset_col.button('Azzera', use_container_width=True):
        for k in ('output', 'rows', 'count', 'input_text', 'source_name'):
            st.session_state.pop(k, None)
        st.rerun()

if clicked:
    if not input_text.strip():
        st.warning('Inserisci un testo oppure carica un documento leggibile.')
    elif not selected:
        st.warning('Seleziona almeno una categoria da anonimizzare.')
    elif nlp is None:
        st.error('Modello non disponibile. Controlla il percorso e gli eventuali errori di caricamento.')
    else:
        try:
            with st.spinner('Analisi delle entità in corso...'):
                spans = find_entities(input_text, nlp)
                result, rows, count = anonymize(input_text, spans, set(selected))
            st.session_state.update(output=result, rows=rows, count=count,
                                    source_name=Path(upload.name).stem if upload else 'testo')
        except Exception as exc:
            st.error(f'Analisi non completata: {type(exc).__name__}: {exc}')

with right:
    st.subheader('Testo anonimizzato')
    st.caption('Clicca Copia e incolla con Ctrl+V')
    # Stesso elemento descrittivo e spazio verticale dell'input a sinistra.
    st.code(st.session_state.output or 'Il testo anonimizzato apparirà qui.',
            language=None, wrap_lines=True, height=330)
    if st.session_state.output:
        basename = st.session_state.get('source_name', 'testo') + '_anonimizzato'
        # Tre pulsanti compatti sulla stessa riga.
        c_txt, c_docx, c_pdf = st.columns(3, gap='small')
        c_txt.download_button('TXT', data=st.session_state.output.encode('utf-8'),
                              file_name=basename + '.txt', mime='text/plain',
                              use_container_width=True)
        try:
            c_docx.download_button('Word', data=make_docx(st.session_state.output),
                                   file_name=basename + '.docx',
                                   mime='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
                                   use_container_width=True)
            c_pdf.download_button('PDF', data=make_pdf(st.session_state.output),
                                  file_name=basename + '.pdf', mime='application/pdf',
                                  use_container_width=True)
        except ImportError as exc:
            st.error(f'Dipendenza mancante per esportare: {exc}. Installa python-docx e reportlab.')
        st.caption('Word e PDF ricreati dal testo; impaginazione originale non garantita.')
if st.session_state.output:
    n_detected = len(st.session_state.rows)
    c1, c2, c3 = st.columns(3)
    c1.metric('Entità rilevate', n_detected)
    c2.metric('Entità anonimizzate', st.session_state.count)
    c3.metric('Categorie rilevate', len({r['Categoria'] for r in st.session_state.rows}))
    with st.expander('Dettaglio entità rilevate', expanded=False):
        st.warning('Questa tabella contiene dati originali: non condividerla insieme al testo anonimizzato.')
        st.dataframe(st.session_state.rows, use_container_width=True, hide_index=True)
    st.caption('Attenzione: il rilevamento automatico può omettere dati personali o segnalare falsi positivi. Controlla sempre il risultato prima di condividerlo.')
